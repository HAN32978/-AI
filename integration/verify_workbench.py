"""隔离合成资料的真实 HTTP / Chroma / 本地 LLM / Streamlit 验证。"""
import argparse
import hashlib
import csv
import io
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import time
import uuid
import requests
from openpyxl import Workbook


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--model-dir", type=Path, required=True)
    args = parser.parse_args()
    repo = Path(__file__).resolve().parents[1]
    runtime = repo / ".runtime" / "workbench-verification" / uuid.uuid4().hex[:12]
    runtime.mkdir(parents=True)
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    base = f"http://127.0.0.1:{port}/api/v1"
    env = os.environ.copy()
    env.update({"EMBEDDING_MODEL_NAME": str(args.model_dir.resolve()), "VECTOR_DB_DIR": str(runtime / "vectors"),
                "UPLOAD_DIR": str(runtime / "upload"), "DOCUMENT_DB_PATH": str(runtime / "documents.sqlite3"),
                "LOG_FILE": str(runtime / "api.log"), "USE_RERANKER": "false", "HF_HUB_OFFLINE": "1",
                "TRANSFORMERS_OFFLINE": "1", "LLM_PROVIDER": "ollama", "OLLAMA_MODEL_NAME": "qwen2.5:0.5b",
                "MAX_UPLOAD_MB": "1", "PYTHONUTF8": "1"})
    checks = []
    def check(label, condition):
        if not condition:
            raise AssertionError(label)
        checks.append(label)
        print("PASS " + label, flush=True)
    def request(method, path, **kwargs):
        response = requests.request(method, base + path, timeout=kwargs.pop("timeout", 20), **kwargs)
        response.raise_for_status()
        return response
    def upload(name, data, **metadata):
        response = request("POST", "/upload/file", files={"file": (name, data)}, data=metadata)
        check("上传接收 " + name, response.status_code in {200, 202})
        doc_id = response.json()["document_id"]
        deadline = time.monotonic() + 90
        while time.monotonic() < deadline:
            doc = request("GET", f"/documents/{doc_id}", params={"project_id": metadata.get("project_id", "default")}).json()
            if doc["status"] not in {"queued", "parsing", "indexing"}:
                check("处理完成 " + name, doc["status"] != "failed")
                return doc
            time.sleep(0.3)
        raise TimeoutError(f"入库超时：{doc_id}")
    log = (runtime / "server.log").open("w", encoding="utf-8")
    process = None
    def start():
        proc = subprocess.Popen([sys.executable, "-m", "uvicorn", "api.main:app", "--host", "127.0.0.1", "--port", str(port)],
                                cwd=repo / "rag_app", env=env, stdout=log, stderr=log,
                                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        deadline = time.monotonic() + 45
        while time.monotonic() < deadline:
            if proc.poll() is not None:
                raise RuntimeError(f"API 启动失败，日志：{runtime}")
            try:
                if requests.get(base + "/system/health", timeout=1).ok:
                    return proc
            except requests.RequestException:
                pass
            time.sleep(0.3)
        proc.terminate()
        raise TimeoutError("启动超时")
    def stop(proc):
        proc.terminate()
        try:
            proc.wait(timeout=15)
        except subprocess.TimeoutExpired:
            proc.kill()
            proc.wait(timeout=5)
    try:
        process = start()
        check("本地模型服务就绪", request("GET", "/system/model_status").json()["status"] == "ready")
        workbook = Workbook()
        workbook.active.append(["演示9月9日检查记录"])
        workbook.active.append(["序号", "部位", "问题描述", "整改要求", "整改状态", "附加字段"])
        for i in range(25):
            workbook.active.append([i + 1, f"合成区{i}", "=演示问题" if i == 0 else f"演示通道问题{i}", "清理材料", "粘贴照片处", f"附加值{i}"])
        workbook.active.cell(row=3, column=3).data_type = "s"  # 原表文本，避免被 Excel 当作未计算公式。
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        data = buffer.getvalue()
        doc = upload("合成检查表.xlsx", data)
        check("真实 Chroma 入库 25 条", doc["record_count"] == 25 and doc["chunk_count"] == 25 and doc["status"] == "ready")
        duplicate = request("POST", "/upload/file", files={"file": ("合成检查表.xlsx", data)}).json()
        check("重复上传复用文档 ID", duplicate["duplicate"] and duplicate["document_id"] == doc["id"])
        first = request("GET", "/records", params={"date": "9月9日", "page_size": 10}).json()
        last = request("GET", "/records", params={"date": "9月9日", "page_size": 10, "page": 3}).json()
        check("完整查询与分页 25 = 10 + 10 + 5", first["total"] == 25 and len(first["items"]) == 10 and len(last["items"]) == 5)
        check("原表空白/照片占位不推断整改状态", all(row["reported_status"] is None for row in first["items"]))
        check("未登记年份不匹配完整年日期", request("GET", "/records", params={"date": "2026-09-09"}).json()["total"] == 0)
        qa = request("POST", "/qa/ask", json={"question": "查询9月9日全部检查问题", "page_size": 200}, timeout=30).json()
        check("自然语言日期查询完整 25 条", qa["intent"] == "完整记录查询" and qa["total"] == 25 and len(qa["records"]) == 25)
        exported = list(csv.reader(io.StringIO(request("GET", "/records/export", params={"date": "9月9日"}).content.decode("utf-8-sig"))))
        check("CSV 导出 25 行、全字段与公式转义", len(exported) == 26 and exported[1][4].startswith("'=") and "附加字段" in exported[1][11])
        def excel_bytes(rows):
            book = Workbook()
            book.active.append(["日期", "问题描述", "整改状态"])
            for row in rows:
                book.active.append(row)
            result = io.BytesIO()
            book.save(result)
            book.close()
            return result.getvalue()
        irregular_data = excel_bytes([[value, "不规整合成问题", "（粘贴照片处）"] for value in
                                     ["9月9日", "9.9", "9/9", "09-09", "九月九日", "2025-09-09", "2026-09-09", "2026年9月", "无法辨认"]] + [["合计", "9", ""]])
        irregular = upload("不规整合成表.xlsx", irregular_data)
        irregular_rows = request("GET", "/records", params={"document_id": irregular["id"], "date": "9月9日"}).json()
        check("真实 HTTP 不规整日期 7 条匹配与 2 条待核对", irregular_rows["total"] == 7 and irregular_rows["unrecognized_date_count"] == 2 and irregular_rows["unresolved_date_count"] == 2)
        check("真实 HTTP 汇总行排除且解析状态显式部分", irregular["record_count"] == 9 and irregular["status"] == "partial")
        check("真实 HTTP 年份分布与占位状态", irregular_rows["year_distribution"] == {"未登记年份": 5, "2025": 1, "2026": 1} and all(r["reported_status"] is None for r in irregular_rows["items"]))
        review_rows = request("GET", "/records", params={"document_id": irregular["id"], "date_quality_filter": "needs_review"}).json()
        check("日期未识别行可以单独查看", review_rows["total"] == 2)
        day9 = upload("每日检查.xlsx", excel_bytes([["9月9日", "第9日独立记录", ""]]))
        day10 = upload("每日检查.xlsx", excel_bytes([["9月10日", "第10日独立记录", ""]]))
        check("真实 HTTP 同名不同日期不会互相顶替", request("GET", f"/documents/{day9['id']}").json()["is_current"] and day10["is_current"])
        revision = upload("每日检查修订.xlsx", excel_bytes([["9月10日", "第10日修订记录", ""]]), replace_document_id=day10["id"])
        check("明确替换才形成版本组", revision["version_group"] == day10["version_group"] and not request("GET", f"/documents/{day10['id']}").json()["is_current"] and request("GET", f"/documents/{day9['id']}").json()["is_current"])
        empty_project = request("POST", "/projects", json={"name": "无资料项目"}).json()["id"]
        no_evidence = request("POST", "/qa/ask", json={"project_id": empty_project, "question": "规范允许的施工偏差是多少？"}).json()
        check("无文档直接拒答且无来源", no_evidence["intent"] == "依据不足" and not no_evidence["sources"])
        bad = requests.post(base + "/upload/file", files={"file": ("坏文件.dwg", b"not-dwg")}, timeout=20)
        check("文件头不符拒绝 422", bad.status_code == 422)
        large = requests.post(base + "/upload/file", files={"file": ("超限.txt", b"x" * (1024 * 1024 + 1))}, timeout=20)
        check("流式大小限制拒绝 413（隔离设置 1MB）", large.status_code == 413)
        dwg = upload("合成头测试.dwg", b"AC1032" + b"\0" * 64)
        rvt = upload("合成头测试.rvt", bytes.fromhex("d0cf11e0a1b11ae1") + b"\0" * 64)
        ifc = upload("合成头测试.ifc", b"ISO-10303-21;\nHEADER;ENDSEC;DATA;ENDSEC;END-ISO-10303-21;")
        check("DWG/RVT/IFC 原件归档不冒称检索成功", all(d["status"] == "stored_only" and d["chunk_count"] == 0 for d in [dwg, rvt, ifc]))
        import fitz
        pdf = fitz.open()
        pdf.new_page().insert_text((72, 72), "Synthetic drawing text: material concrete C30")
        pdf_bytes = pdf.tobytes()
        pdf.close()
        export = upload("合成导出说明.pdf", pdf_bytes, source_document_id=rvt["id"], category="复查依据")
        check("PDF 真实文字解析并关联 RVT 原件", export["status"] == "ready" and export["source_document_id"] == rvt["id"])
        import ezdxf
        drawing = ezdxf.new("R2010")
        drawing.modelspace().add_text("Synthetic material C30")
        drawing.modelspace().add_mtext("Synthetic drawing explanation")
        path = runtime / "drawing.dxf"
        drawing.saveas(path)
        dxf = upload("合成文字图纸.dxf", path.read_bytes(), source_document_id=dwg["id"])
        check("DXF 真实文字解析标记部分解析", dxf["status"] == "partial" and dxf["chunk_count"] >= 2)
        check("归档原件可下载", request("GET", f"/documents/{dwg['id']}/download").content == b"AC1032" + b"\0" * 64)
        other = request("POST", "/projects", json={"name": "隔离项目B"}).json()["id"]
        doc_b = upload("项目B.txt", "本项目测试口令是蓝色。".encode(), project_id=other)
        doc_a = upload("项目A.txt", "本项目测试口令是橙色。".encode(), category="整改回复")
        scoped = request("POST", "/qa/ask", json={"question": "本项目测试口令是什么？", "query_mode": "semantic", "document_id": doc_a["id"]}, timeout=120).json()
        check("真实本地模型引用限定文档", bool(scoped["answer"].strip()) and bool(scoped["sources"]) and all(s["document_id"] == doc_a["id"] for s in scoped["sources"]))
        cross = requests.get(base + "/documents/" + doc_a["id"], params={"project_id": other}, timeout=10)
        check("跨项目文档查询拒绝 422", cross.status_code == 422)
        history_b = request("GET", "/qa/history", params={"project_id": other}).json()
        check("相同会话 ID 在另一项目无历史", history_b["history"] == [])
        issue = request("POST", "/issues", json={"record_id": first["items"][0]["id"], "actor": "演示登记人"}).json()
        issue_id = issue["id"]
        def event(action, attachments, actor="演示复查人"):
            return requests.post(base + f"/issues/{issue_id}/events", json={"action": action, "actor": actor,
                                 "note": "合成测试处理说明", "attachment_ids": attachments, "evidence_date": "2026-09-11"}, timeout=15)
        check("未回复直接销项被拒绝", event("recheck_pass", [doc_a["id"]]).status_code == 422)
        check("回复缺附件被拒绝", event("reply", []).status_code == 422)
        check("跨项目附件被拒绝", event("reply", [doc_b["id"]]).status_code == 422)
        check("有效整改回复", event("reply", [doc_a["id"]]).json()["status"] == "reply_received")
        check("人工申请复查", event("request_recheck", []).json()["status"] == "recheck_pending")
        check("人工复查带依据后销项", event("recheck_pass", [export["id"]]).json()["status"] == "closed")
        before = request("GET", "/documents").json()["items"]
        stop(process)
        process = start()
        persisted = request("GET", f"/issues/{issue_id}").json()
        check("API 重启后人工销项及 4 条事件保留", persisted["status"] == "closed" and len(persisted["events"]) == 4)
        check("API 重启不把 UUID 原件重复登记", len(request("GET", "/documents").json()["items"]) == len(before))
        check("历史图纸接口主应用已移除", requests.post(base + "/review/run", json={}, timeout=10).status_code == 404)
        check("旧版单独清空向量被停用", requests.post(base + "/knowledge/clear", timeout=10).status_code == 409)
        stop(process)
        env.update({"LLM_PROVIDER": "openai", "OPENAI_API_KEY": ""})
        process = start()
        missing_key = requests.post(base + "/qa/ask", json={"question": "测试口令？", "query_mode": "semantic", "document_id": doc_a["id"]}, timeout=90)
        check("真实 HTTP 缺少 API Key 返回 503", missing_key.status_code == 503)
        still_no_evidence = request("POST", "/qa/ask", json={"project_id": empty_project, "question": "施工规范要求？"}).json()
        check("模型未配置时无依据范围仍可拒答", still_no_evidence["intent"] == "依据不足")
        stop(process)
        env.update({"LLM_PROVIDER": "ollama"})
        process = start()
        # 使用本机已安装 Streamlit 1.39 的 AppTest，实际连接上面的隔离 API。
        os.environ["RAG_API_BASE"] = base
        sys.path.insert(0, str(repo / "rag_app" / "frontend"))
        from streamlit.testing.v1 import AppTest
        pages = [repo / "rag_app" / "frontend" / "app.py"] + sorted((repo / "rag_app" / "frontend" / "pages").glob("*.py"))
        for page in pages:
            app = AppTest.from_file(str(pages[0]), default_timeout=40).run()
            if page != pages[0]:
                app.switch_page("pages/" + page.name).run()
            check("Streamlit 页面无运行异常 " + page.name, not app.exception)
        stop(process)
        process = None
        log.flush()
        server_log = (runtime / "server.log").read_bytes()
        (repo / "docs" / "工作台集成服务日志_2026-10-08.txt").write_bytes(server_log)
        report = {"status": "passed", "checked_at": time.strftime("%Y-%m-%d %H:%M:%S"), "check_count": len(checks),
                  "checks": checks, "structured_total": 25, "semantic_source_count": len(scoped["sources"]),
                  "server_log_sha256": hashlib.sha256(server_log).hexdigest(),
                  "scope": "隔离规则及不规整合成资料、真实 HTTP/Chroma/本地 Ollama/Streamlit AppTest；DWG/RVT/IFC 仅验证基础文件头与归档，未验证真实模型/代理对象解析；未评估工程规范准确率。"}
        (repo / "docs" / "工作台回归验证.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({k: v for k, v in report.items() if k != "checks"}, ensure_ascii=False), flush=True)
    finally:
        if process and process.poll() is None:
            stop(process)
        log.close()


if __name__ == "__main__":
    main()
