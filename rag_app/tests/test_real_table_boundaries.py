"""不规整表格与审查报告逐项对应的真实 SQLite/Excel 回归。"""
import asyncio
import sqlite3
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import Mock, patch
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_community.vectorstores import FAISS
from fastapi import HTTPException
from core.date_parser import parse_business_date
from core.file_ingestion import read_excel_records, validate_file
from core.rag_chain import RAGChain
from core.memory_manager import MemoryManger
from core.vector_store import VectorStoreManager
from core.intent_recognizer import IntentRecognizer
from api.dependencies import LazyRAGChain
from api.routes.qa import ask_question, QuestionRequest
import test_catalog_workflow as fixtures


class BoundaryTests(unittest.TestCase):
    setUp = fixtures.CatalogTests.setUp
    tearDown = fixtures.CatalogTests.tearDown
    workbook = fixtures.CatalogTests.workbook
    register_ready = fixtures.CatalogTests.register_ready
    def irregular(self, rows, headers=None, name="不规整.xlsx"):
        from openpyxl import Workbook
        path = self.folder / name
        book = Workbook()
        sheet = book.active
        sheet.append(headers or ["日期", "问题描述", "整改状态"])
        for row in rows:
            sheet.append(row)
        book.save(path)
        book.close()
        return path

    def test_same_filename_different_days_both_remain_current(self):
        first = self.register_ready(self.irregular([["9月9日", "演示问题9", ""]], name="检查记录.xlsx"))
        second = self.register_ready(self.irregular([["9月10日", "演示问题10", ""]], name="检查记录.xlsx"))
        self.assertTrue(self.catalog.document(first["id"])["is_current"])
        self.assertTrue(self.catalog.document(second["id"])["is_current"])
        self.assertEqual(self.catalog.query_records(date="9月9日")["total"], 1)
        self.assertEqual(self.catalog.query_records(date="9月10日")["total"], 1)

    def test_common_dates_all_match_without_guessing_year_or_day(self):
        values = ["9月9日", "9.9", "9/9", "09-09", "九月九日", "二〇二六年九月九日", "2026-09-09 00:00:00", "2026年9月", "无法辨认"]
        self.register_ready(self.irregular([[value, "合成问题", ""] for value in values]))
        data = self.catalog.query_records(date="9月9日")
        self.assertEqual(data["total"], 7)
        self.assertEqual(data["unresolved_date_count"], 2)
        self.assertEqual(data["unrecognized_date_count"], 2)
        self.assertEqual(data["completeness"], "date_unresolved")
        self.assertEqual(self.catalog.query_records(date_quality_filter="needs_review")["total"], 2)
        self.assertEqual(self.catalog.query_records()["total"], 9)
        with self.assertRaises(ValueError):
            parse_business_date("2026年9月")

    def test_summary_rows_are_reported_and_not_problem_records(self):
        doc = self.register_ready(self.irregular([["9月9日", "演示问题", ""], ["合计", "1", ""], ["小计", "1", ""]]))
        import json
        report = json.loads(doc["parse_report"])
        self.assertEqual(doc["record_count"], 1)
        self.assertEqual(len(report["summary_rows"]), 2)
        self.assertEqual(report["summary_rows"][0]["row"], 3)

    def test_status_placeholders_and_real_state_with_photo_note(self):
        statuses = ["粘贴照片处。", "（粘贴照片处）", "图片占位!", "已整改（照片待补充）"]
        self.register_ready(self.irregular([["9月9日", "演示问题", status] for status in statuses]))
        rows = self.catalog.query_records()["items"]
        self.assertTrue(all(r["reported_status"] is None for r in rows[:3]))
        self.assertEqual(rows[3]["reported_status"], statuses[3])
        self.assertEqual(rows[0]["cells"]["整改状态"], statuses[0])

    def test_duplicate_columns_keep_values_and_conflicting_dates_are_flagged(self):
        path = self.irregular([["9月9日", "9月10日", "一号问题", "二号问题"]], ["日期", "日期", "问题描述", "问题描述"])
        self.register_ready(path)
        data = self.catalog.query_records()
        row = data["items"][0]
        self.assertEqual(row["cells"]["问题描述#2"], "二号问题")
        self.assertEqual(row["cells"]["日期#2"], "9月10日")
        self.assertEqual(data["unrecognized_date_count"], 1)
        self.assertEqual(self.catalog.query_records(date="9月9日")["total"], 0)

    def test_unrecognized_header_preserves_first_row(self):
        path = self.irregular([["甲", "乙"]], headers=["不是标准表头", "说明正文"])
        report = {}
        rows = read_excel_records(path, "sample", report=report)
        self.assertEqual(len(rows), 2)
        self.assertTrue(report["header_warnings"])

    def test_month_day_query_exposes_year_distribution(self):
        self.register_ready(self.irregular([["2025-09-09", "问题A", ""], ["2026-09-09", "问题B", ""], ["9月9日", "问题C", ""]]))
        data = self.catalog.query_records(date="9月9日")
        self.assertEqual(data["year_distribution"], {"未登记年份": 1, "2025": 1, "2026": 1})
        self.assertIn("跨年份", data["note"])
        self.assertEqual(self.catalog.query_records(date="2026-09-09")["total"], 1)

    def test_empty_scope_returns_without_model_initialization(self):
        with patch("api.routes.qa.get_catalog", return_value=self.catalog), patch("api.dependencies.get_rag_chain", side_effect=AssertionError("不得初始化模型")):
            result = asyncio.run(ask_question(QuestionRequest(question="规范允许的偏差是多少？"), LazyRAGChain()))
        self.assertEqual(result.intent, "依据不足")

    def test_empty_retrieval_returns_without_llm_even_with_history(self):
        chain = RAGChain.__new__(RAGChain)
        chain.retriever = Mock()
        chain.retriever.retrieve.return_value = []
        chain.memory_manager = MemoryManger()
        chain.memory_manager.add_exchange("history", "施工问题", "旧回答")
        chain.intent_recognizer = IntentRecognizer()
        chain.llm = None
        with patch("core.rag_chain.get_llm", side_effect=AssertionError("不得调用 LLM")), patch("core.rag_chain.ensure_llm_ready", side_effect=AssertionError("不得查询模型服务")):
            result = chain.ask("它的规范要求是多少？", "history", ["example-id"])
        self.assertEqual(result["intent"], "依据不足")

    def test_model_configuration_value_error_becomes_503(self):
        self.register_ready(self.workbook(count=1))
        with patch("api.routes.qa.get_catalog", return_value=self.catalog), patch("api.dependencies.get_rag_chain", side_effect=ValueError("缺少 API Key")):
            with self.assertRaises(HTTPException) as caught:
                asyncio.run(ask_question(QuestionRequest(question="规范要求", query_mode="semantic"), LazyRAGChain()))
        self.assertEqual(caught.exception.status_code, 503)

    def test_source_table_cannot_serve_as_closure_evidence(self):
        doc = self.register_ready(self.workbook(count=1), category="整改回复")
        record = self.catalog.query_records()["items"][0]
        issue = self.catalog.create_issue("default", "", "登记人", record_id=record["id"])
        with self.assertRaisesRegex(ValueError, "问题原表"):
            self.catalog.add_issue_event(issue["id"], "reply", "回复人", "说明", [doc["id"]], "2026-09-10")

    def test_recheck_cannot_precede_reply_and_requires_category(self):
        path = self.folder / "reply.txt"
        path.write_text("真实测试回复", encoding="utf-8")
        reply = self.register_ready(path, category="整改回复")
        issue = self.catalog.create_issue("default", "问题", "登记人")
        self.catalog.add_issue_event(issue["id"], "reply", "回复人", "回复说明", [reply["id"]], "2026-09-11")
        self.catalog.add_issue_event(issue["id"], "request_recheck", "申请人", "申请复查")
        with self.assertRaisesRegex(ValueError, "不能早于"):
            self.catalog.add_issue_event(issue["id"], "recheck_pass", "复查人", "通过", [reply["id"]], "2026-09-10")
        with self.assertRaisesRegex(ValueError, "复查依据"):
            self.catalog.add_issue_event(issue["id"], "recheck_pass", "复查人", "通过", [reply["id"]], "2026-09-12")

    def test_misleading_dxf_header_rejected(self):
        path = self.folder / "false.dxf"
        path.write_bytes(b"hello SECTION 0")
        with self.assertRaises(ValueError):
            validate_file(path)

    def test_legacy_schema_migration_restores_hidden_records_and_reparses(self):
        from core.document_catalog import DocumentCatalog
        from core.file_ingestion import IngestionWorker
        first = self.register_ready(self.irregular([["9月9日", "旧问题9", ""]], name="检查记录.xlsx"))
        # 原件是不可变归档，模拟第二份同名原件需要另一条存储路径。
        first_path = self.folder / "archived-first.xlsx"
        first_path.write_bytes(Path(first["stored_path"]).read_bytes())
        with self.catalog.connect() as db:
            db.execute("UPDATE documents SET stored_path=? WHERE id=?", (str(first_path), first["id"]))
        second = self.register_ready(self.irregular([["9月10日", "旧问题10", ""]], name="检查记录.xlsx"))
        with self.catalog.connect() as db:
            db.execute("UPDATE documents SET is_current=0 WHERE id=?", (first["id"],))
            for table, column in [("documents", "version_group"), ("documents", "parse_report"), ("records", "date_quality"), ("records", "date_error"), ("issue_events", "evidence_date")]:
                db.execute(f"ALTER TABLE {table} DROP COLUMN {column}")
        migrated = DocumentCatalog(self.catalog.path)
        self.assertTrue(migrated.document(first["id"])["is_current"])
        self.assertEqual(migrated.document(second["id"])["status"], "queued")
        worker = IngestionWorker(migrated, lambda: self.vectors)
        while worker.process_next():
            pass
        self.assertEqual(migrated.query_records(date="9月9日")["total"], 1)
        self.assertEqual(migrated.query_records(date="9月10日")["total"], 1)
        after = DocumentCatalog(self.catalog.path)
        self.assertEqual(after.document(first["id"])["status"], "ready")


class ToyEmbeddings(Embeddings):
    def embed_documents(self, texts):
        return [[float("A" in text), float("B" in text), 1.0] for text in texts]
    def embed_query(self, text):
        return self.embed_documents([text])[0]


class FaissFilterTests(unittest.TestCase):
    def test_real_faiss_filter_applies_to_both_search_methods(self):
        store = VectorStoreManager.__new__(VectorStoreManager)
        store._store = FAISS.from_documents([Document(page_content="A", metadata={"document_id": "a"}),
                                            Document(page_content="B", metadata={"document_id": "b"})], ToyEmbeddings())
        with patch("core.vector_store.settings.VECTOR_STORE_TYPE", "faiss"):
            docs = store.similarity_search("B", k=1, filter_dict={"document_id": {"$in": ["a"]}})
            scores = store.similarity_search_with_score("B", k=1, filter_dict={"document_id": {"$in": ["a"]}})
            from core.retriever import RAGRetriever
            retriever = RAGRetriever.__new__(RAGRetriever)
            retriever.vector_store = store
            retriever.reranker = None
            chain_docs = retriever.get_compresstion_retriever({"k": 1, "filter": {"document_id": {"$in": ["a"]}}}).invoke("B")
        self.assertEqual([d.metadata["document_id"] for d in docs], ["a"])
        self.assertEqual([d.metadata["document_id"] for d, _ in scores], ["a"])
        self.assertEqual([d.metadata["document_id"] for d in chain_docs], ["a"])
