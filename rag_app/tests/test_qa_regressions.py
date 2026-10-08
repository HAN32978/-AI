"""针对实际发生的 Excel 403 和模型缺失错误进行回归验证。"""
import asyncio
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch, Mock

from fastapi import HTTPException
from openpyxl import Workbook
from api.routes.qa import ask_question, QuestionRequest
from core.document_loader import DocumentLoader
from core.llm_client import ensure_llm_ready, LLMServiceError


class QARegressionTests(unittest.TestCase):
    def test_xlsx_without_network_preserves_sheet_row_and_cells(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sample.xlsx"
            workbook = Workbook()
            workbook.active.title = "检查记录"
            workbook.active.append(["日期", "问题", "整改要求"])
            workbook.active.append(["9月9日", "演示通道堆放材料", "演示要求：清理"])
            workbook.create_sheet("空表")
            workbook.save(path)
            workbook.close()
            with patch("socket.create_connection", side_effect=AssertionError("解析 Excel 不得联网")):
                docs = DocumentLoader().load_file(str(path))
            self.assertTrue(docs)
            row_docs = [doc for doc in docs if doc.metadata.get("row") == 2]
            self.assertTrue(row_docs)
            self.assertIn("日期：9月9日", row_docs[0].page_content)
            self.assertIn("问题：演示通道堆放材料", row_docs[0].page_content)
            self.assertEqual(row_docs[0].metadata["sheet_name"], "检查记录")

    def test_missing_model_has_actionable_message(self):
        response = Mock()
        response.json.return_value = {"models": []}
        with patch("core.llm_client.settings.LLM_PROVIDER", "ollama"), patch("core.llm_client.requests.get", return_value=response):
            with self.assertRaisesRegex(LLMServiceError, "ollama pull"):
                ensure_llm_ready()

    def test_merged_title_is_not_a_header_or_data_row(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "title.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["演示9月9日验收记录"])
            sheet.merge_cells("A1:C1")
            sheet.append(["序号", "整改问题描述", "整改状态"])
            sheet.append([1, "演示通道堆放材料", "待整改"])
            workbook.save(path)
            workbook.close()
            docs = DocumentLoader().load_file(str(path))
            self.assertEqual(len(docs), 1)
            self.assertIn("整改问题描述：演示通道堆放材料", docs[0].page_content)
            self.assertEqual(docs[0].metadata["row"], 3)

    def test_failed_qa_returns_503_instead_of_success_answer(self):
        chain = Mock()
        chain.ask.side_effect = LLMServiceError("模型未安装")
        with self.assertRaises(HTTPException) as result:
            asyncio.run(ask_question(QuestionRequest(question="演示问题"), chain))
        self.assertEqual(result.exception.status_code, 503)
        self.assertEqual(result.exception.detail, "模型未安装")

    def test_excel_source_without_page_accepts_sheet_and_row(self):
        chain = Mock()
        chain.ask.return_value = {"answer": "演示回答", "intent": "通用问答", "sources": [
            {"source": "sample.xlsx", "page": None, "sheet_name": "检查记录", "row": 3, "content": "演示记录"}
        ]}
        result = asyncio.run(ask_question(QuestionRequest(question="演示问题"), chain))
        self.assertIsNone(result.sources[0].page)
        self.assertEqual(result.sources[0].row, 3)


if __name__ == "__main__":
    unittest.main()
